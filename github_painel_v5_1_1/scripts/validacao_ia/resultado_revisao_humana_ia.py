"""Resultado da revisao humana independente da aba Inteligencia artificial (50 casos).
Executar a partir de Processado/:  python validacao_ia/resultado_revisao_humana_ia.py
Entradas (somente leitura): validacao_ia/revisao_humana_ia.xlsx (coluna H preenchida pela revisora; so as 50 linhas do
  grupo 'Classes da aba IA'); validacao_ia/revisao_humana_ia_julgamentos_anteriores.csv (verificacao por amostragem).
Regras: concordancia da classe = linhas 'certo' / linhas revisadas; classe validada com pelo menos 80%;
  concordancia entre revisoes = linhas em que a revisao humana ('certo'/'errado') coincide com o julgamento anterior ('sim'/'nao').
  Motivos dos erros apontados na revisao humana registrados em MOTIVO_REVISAO; o dicionario nao muda.
  Verificacao por valor: as 10 compras de maior valor da classe Capacitacao (grupo 'Capacitação: maiores valores', valores de
  validacao_ia/capacitacao_maiores_valores_revisadas.csv, lista revisada e congelada): compras certas e parcela do valor nas compras marcadas como erradas.
Saidas: validacao_ia/resultado_revisao_humana_ia.csv (resultado por classe, concordancia geral e verificacao por valor);
  validacao_ia/divergencias_revisao_humana_ia.csv (casos em que as duas revisoes discordam ou a revisao humana marcou errado)
"""
import csv

from openpyxl import load_workbook

V = "validacao_ia/"
CLASSES = ["Capacitação", "Acréscimo à licença", "Solução de IA (exclusiva ou em pacote)", "Outros"]
MOTIVO_REVISAO = {  # erros apontados na revisao humana
    "09002106000772025": "falso positivo: compra de alimentação (lanches para curso de IA), não capacitação",
    "08001207000032023": "palestra sobre ChatGPT: deveria ser Capacitação",
    "15813705903072025": "manutenção de sistema de CFTV: não é IA",
}
ws = load_workbook(V + "revisao_humana_ia.xlsx", read_only=True)["revisao"]
todas = list(ws.iter_rows(min_row=2, values_only=True))
linhas = [r for r in todas if r[0] == "Classes da aba IA"]
extra = [r for r in todas if r[0] == "Capacitação: maiores valores"]
assert len(linhas) == 50 and all(str(r[7]).strip().lower() in ("certo", "errado") for r in linhas), "revisão incompleta"
ant = {r["id"]: r for r in csv.DictReader(open(V + "revisao_humana_ia_julgamentos_anteriores.csv", encoding="utf-8"), delimiter=";")}
res, div = {}, []
for r in linhas:
    c, h = r[3], str(r[7]).strip().lower()
    a = ant[r[1]]
    d = res.setdefault(c, [0, 0, 0])
    d[0] += 1
    d[1] += h == "certo"
    conc = (h == "certo") == (a["pertinente_julgamento_anterior"] == "sim")
    d[2] += conc
    if not conc or h == "errado":
        mot = MOTIVO_REVISAO.get(r[1]) or (f"revisão humana marcou certo; julgamento anterior (não pertinente): {a['motivo']}" if h == "certo" else a["motivo"])
        div.append([r[1], c, h, a["pertinente_julgamento_anterior"], "sim" if conc else "não", mot])
assert set(MOTIVO_REVISAO) <= {x[0] for x in div}
with open(V + "resultado_revisao_humana_ia.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.writer(fo, delimiter=";")
    w.writerow(["classe", "casos_revisados", "certo", "errado", "concordancia_pct", "validada", "concordancia_entre_revisoes", "concordancia_entre_revisoes_pct",
                "valor_revisado", "valor_em_compras_erradas", "pct_do_valor_em_compras_erradas"])
    for c in CLASSES:
        n, p, k = res[c]
        w.writerow([c, n, p, n - p, round(100 * p / n, 1), "sim" if 100 * p / n >= 80 else "não", k, round(100 * k / n, 1)])
    n, p, k = (sum(x[i] for x in res.values()) for i in range(3))
    w.writerow(["Total", n, p, n - p, round(100 * p / n, 1), "", k, round(100 * k / n, 1)])
    # verificacao por valor (10 maiores compras de Capacitacao)
    if extra:
        assert len(extra) == 10 and all(str(r[7]).strip().lower() in ("certo", "errado") for r in extra), "verificação por valor incompleta"
        mv = {r["id_compra"]: float(r["valor_nominal"]) for r in csv.DictReader(open(V + "capacitacao_maiores_valores_revisadas.csv", encoding="utf-8"), delimiter=";")}
        vt = sum(mv[r[1]] for r in extra)
        ve = sum(mv[r[1]] for r in extra if str(r[7]).strip().lower() == "errado")
        pc = sum(str(r[7]).strip().lower() == "certo" for r in extra)
        w.writerow(["Capacitação: verificação por valor (10 maiores compras)", len(extra), pc, len(extra) - pc, round(100 * pc / len(extra), 1), "não se aplica", "", "",
                    round(vt, 2), round(ve, 2), round(100 * ve / vt, 1)])
        print("verificação por valor:", pc, "de", len(extra), "| valor revisado", round(vt / 1e6, 2), "mi | errado", round(ve / 1e6, 2), "mi")
with open(V + "divergencias_revisao_humana_ia.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.writer(fo, delimiter=";")
    w.writerow(["id_compra", "classe", "revisao_humana", "julgamento_anterior", "revisoes_concordam", "motivo"])
    w.writerows(div)
for c in CLASSES:
    n, p, k = res[c]
    print(c, f"{p}/{n}", f"{100 * p / n:.0f}%", "| concordância entre revisões", k)
print("concordância geral:", sum(x[2] for x in res.values()), "de", sum(x[0] for x in res.values()))
for x in div:
    print(x)
