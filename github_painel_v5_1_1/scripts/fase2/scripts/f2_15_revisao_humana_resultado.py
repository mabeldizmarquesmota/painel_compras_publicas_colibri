"""Fase 2, etapa 15: resultado da revisao humana independente e pendencias do dicionario das categorias.
Executar a partir de Processado/, depois do preenchimento de fase2/validacao/revisao_humana.xlsx:
  python fase2/scripts/f2_15_revisao_humana_resultado.py
Entradas: fase2/validacao/revisao_humana.xlsx (coluna 'revisão (certo/errado)' preenchida pela revisora);
  fase2/validacao/revisao_humana_julgamentos_anteriores.csv (julgamentos da verificacao por amostragem, mesmas linhas);
  fase2/f2.duckdb (somente leitura: compra, objeto e forma de classificacao de cada item); dicionario_termos.csv.
Regras:
  - concordancia da classe = linhas 'certo' / linhas revisadas; classe validada quando >= 80%;
  - concordancia entre revisoes = linhas em que a revisao humana ('certo'/'errado') coincide com o julgamento
    anterior ('sim'/'nao');
  - pendencias: itens da amostra cujo texto esta fora da tecnologia da categoria; os casos apontados na revisao humana
    (lista CASOS_REVISAO) e, separadamente, outros casos evidentes da mesma amostra (lista OUTROS_CASOS). O termo
    provavel e o termo de inclusao do dicionario das categorias encontrado no objeto (compra classificada pelo objeto)
    (objeto e informacao complementar, campos da classificacao) ou no texto do item (compra classificada pelos itens). O dicionario nao e alterado.
Saidas: fase2/validacao/resultado_revisao_humana.csv; fase2/validacao/concordancia_entre_revisoes.csv;
        fase2/pendencias_dicionario_categorias.md
"""
import csv
import re
import sys
import unicodedata

import duckdb
from openpyxl import load_workbook

sys.path.insert(0, "scripts_painel_v2")
from formatacao import fmt_pct, inteiro  # noqa: E402

V = "fase2/validacao/"
ws = load_workbook(V + "revisao_humana.xlsx")["revisao"]
linhas = [r for r in ws.iter_rows(min_row=2, values_only=True)]
assert len(linhas) == 120 and all(str(r[7]).strip().lower() in ("certo", "errado") for r in linhas), "revisão incompleta"
ant = list(csv.DictReader(open(V + "revisao_humana_julgamentos_anteriores.csv", encoding="utf-8"), delimiter=";"))
assert [a["id"] for a in ant] == [r[1] for r in linhas]


def classe(r):
    return r[3] if r[0] == "Modelo de aquisição" else r[0]


res = {}
for r, a in zip(linhas, ant):
    k = (r[0], classe(r))
    d = res.setdefault(k, {"n": 0, "certo": 0, "concorda": 0})
    d["n"] += 1
    certo = str(r[7]).strip().lower() == "certo"
    d["certo"] += certo
    d["concorda"] += certo == (a["pertinente_julgamento_anterior"] == "sim")
with open(V + "resultado_revisao_humana.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.writer(fo, delimiter=";")
    w.writerow(["grupo", "classe", "linhas_revisadas", "certo", "errado", "concordancia_pct", "validada", "concordancia_entre_revisoes_pct"])
    for (g, c), d in res.items():
        p = 100 * d["certo"] / d["n"]
        w.writerow([g, c, d["n"], d["certo"], d["n"] - d["certo"], round(p, 1), "sim" if p >= 80 else "não", round(100 * d["concorda"] / d["n"], 1)])
tot_conc = sum(d["concorda"] for d in res.values())
with open(V + "concordancia_entre_revisoes.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.writer(fo, delimiter=";")
    w.writerow(["grupo", "id", "classe", "revisao_humana", "julgamento_anterior", "concordam"])
    for r, a in zip(linhas, ant):
        rh = str(r[7]).strip().lower()
        w.writerow([r[0], r[1], classe(r), rh, a["pertinente_julgamento_anterior"], "sim" if (rh == "certo") == (a["pertinente_julgamento_anterior"] == "sim") else "não"])
for (g, c), d in res.items():
    print(g, "|", c, "|", d["n"], d["certo"], fmt_pct(100 * d["certo"] / d["n"]), "| entre revisões:", fmt_pct(100 * d["concorda"] / d["n"]))
print("concordância geral entre revisões:", tot_conc, "de", len(linhas), fmt_pct(100 * tot_conc / len(linhas)))

# ---------------------------------------------------------------- pendencias do dicionario das categorias
CASOS_REVISAO = {  # palavra-chave do texto do item -> descricao (casos apontados na revisao humana)
    "cateter": "cateter (material médico)", "oftalmolog": "equipamento de uso oftalmológico", "polpa de fruta": "polpa de fruta (gênero alimentício)",
    "economico / financeira": "consultoria econômico-financeira", "norma din": "norma técnica (acesso a normas) classificada como PaaS"}
OUTROS_CASOS = {"plug audio": "plugue de áudio", "regua de sinalizacao": "régua de sinalização de trânsito", "escala de desenho": "escala de desenho de arquitetura",
                "camara escura": "câmara escura de fotografia", "insignia": "insígnia (serviço de desinsetização)", "bateria estacionaria": "bateria estacionária de nobreak",
                "cooperacao / convenios": "acordo de cooperação técnica"}
norm = lambda t: re.sub(r"\s+", " ", unicodedata.normalize("NFKD", (t or "").replace("\xa0", " ")).encode("ascii", "ignore").decode().lower()).strip()
dic = [r for r in csv.DictReader(open("dicionario_termos.csv", encoding="utf-8")) if r["tipo"] == "inclusao"]
con = duckdb.connect("fase2/f2.duckdb", read_only=True)
con.sql("ATTACH 'scripts_parte0/p0.duckdb' AS p (READ_ONLY)")
NOME = {"nuvem": "Nuvem", "plataformas": "Plataformas e licenciamento", "data_center": "Data center", "seguranca": "Segurança cibernética",
        "ia": "Inteligência artificial", "identidade": "Identidade e infraestrutura pública digital"}
out = {"revisao": [], "outros": []}
for r in linhas:
    if r[0] != "Modelo de aquisição":
        continue
    t = norm(r[5])
    for lista, casos in (("revisao", CASOS_REVISAO), ("outros", OUTROS_CASOS)):
        hit = next((d for k, d in casos.items() if k in t), None)
        if not hit:
            continue
        item = r[1]
        segs = con.sql(f"""SELECT e.segmento, e.id_compra, c.local_match, any_value(concat_ws(' | ', b.objeto_compra, b.informacao_complementar)), any_value(concat_ws(' | ', i.descricao, i.descricao_resumida, i.descricao_detalhada))
            FROM itens_escopo e JOIN p.cat_compra c ON c.id_compra = e.id_compra AND c.segmento = e.segmento
            JOIN compras_escopo b ON b.id_compra = e.id_compra AND b.segmento = e.segmento JOIN itens_escopo i ON i.id_compra_item = e.id_compra_item AND i.segmento = e.segmento
            WHERE e.id_compra_item = '{item}' GROUP BY 1, 2, 3""").fetchall()
        for seg, compra, local, ob, txt in segs:
            onde = norm(ob) if local in ("objeto", "ambos") else norm(txt)
            termos = sorted({d["termo_regex"] for d in dic if d["segmento"] == seg and re.search(d["termo_regex"], onde)})
            out[lista].append([compra, item, NOME[seg], hit, "objeto ou informação complementar da compra" if local in ("objeto", "ambos") else "texto do item",
                               "; ".join(termos) or "(termo não localizado)", r[3]])
md = ["# Pendências do dicionário das categorias", "",
      "Itens da amostra de revisão humana (120 casos) cujo texto está fora da tecnologia da categoria em que a compra foi classificada. "
      "Termo provável: termo de inclusão do dicionário da categoria (dicionario_termos.csv) encontrado onde a compra foi classificada "
      "(no objeto, quando a compra inteira entra na categoria; no texto do item, quando só os itens entram). O dicionário não foi alterado; "
      "estes casos ficam para a próxima revisão do dicionário. Gerado por `fase2/scripts/f2_15_revisao_humana_resultado.py`.", ""]
for lista, tit in (("revisao", "Casos apontados na revisão humana"), ("outros", "Outros casos evidentes na mesma amostra (a confirmar)")):
    md += [f"## {tit}", "", "| Compra (id) | Item (id) | Categoria | Item fora da tecnologia | Classificação pela | Termo provável | Classe de modelo de aquisição |",
           "|---|---|---|---|---|---|---|"] + ["| " + " | ".join(x) + " |" for x in out[lista]] + [""]
md += ["Nas compras classificadas pelo objeto, a compra inteira entra na categoria, inclusive itens alheios à tecnologia; "
       "o ajuste possível é restringir o escopo de itens nessas compras ou acrescentar termos de exclusão. Nenhuma das duas mudanças foi feita.", ""]
open("fase2/pendencias_dicionario_categorias.md", "w", encoding="utf-8").write("\n".join(md))
print(len(out["revisao"]), "casos da revisão;", len(out["outros"]), "outros casos")
