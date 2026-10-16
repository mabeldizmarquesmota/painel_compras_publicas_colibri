"""Fase 2, etapa 14: planilha para revisao humana de uma subamostra dos itens e compras ja julgados.
Executar a partir de Processado/:  python fase2/scripts/f2_14_revisao_humana.py
Entradas: fichas e julgamentos de fase2/validacao/ (rodada final de cada classe aprovada: classe 1 = a1; 2, 5 e 6 = a2;
  3 e 7 = a3; marcas e empresas publicas de TI = f2_marca_a1); fase2/f2.duckdb (categorias e vencedores de cada item).
Regra: sorteio por md5(id || 'f2_revisao_humana'): 15 itens por classe aprovada, 20 compras com marca e 10 com empresa
  publica de TI, entre os ja julgados. A planilha nao traz os julgamentos anteriores (revisao independente); a coluna
  'revisão (certo/errado)' fica vazia. Os julgamentos anteriores das mesmas linhas ficam em arquivo separado, para comparacao.
Saidas: fase2/validacao/revisao_humana.xlsx; fase2/validacao/revisao_humana_julgamentos_anteriores.csv
"""
import csv
import hashlib

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SEM = "f2_revisao_humana"
V = "fase2/validacao/"
h = lambda x: hashlib.md5((x + SEM).encode()).hexdigest()
ler = lambda n: list(csv.DictReader(open(V + n, encoding="utf-8"), delimiter=";"))
NOME = {"nuvem": "Nuvem", "plataformas": "Plataformas e licenciamento", "data_center": "Data center", "seguranca": "Segurança cibernética",
        "ia": "Inteligência artificial", "identidade": "Identidade e infraestrutura pública digital"}
FINAL = {1: "a1", 2: "a2", 3: "a3", 5: "a2", 6: "a2", 7: "a3"}
con = duckdb.connect("fase2/f2.duckdb", read_only=True)
cats_item = dict(con.sql("SELECT id_compra_item, string_agg(DISTINCT segmento, ',') FROM itens_escopo GROUP BY 1").fetchall())
venc_item = dict(con.sql("SELECT id_compra_item, string_agg(DISTINCT nome_fornecedor, ' / ') FROM res_escopo GROUP BY 1").fetchall())
cats_compra = dict(con.sql("SELECT id_compra, string_agg(DISTINCT segmento, ',') FROM compras_escopo GROUP BY 1").fetchall())
cat_txt = lambda s: "; ".join(NOME[c] for c in (s or "").split(",") if c)

linhas, ant = [], []
for o, rod in FINAL.items():
    fichas = {r["id_compra_item"]: r for r in ler(f"fichas_modelo_f2_modelo_{rod}.csv") if int(r["ordem"]) == o}
    jul = {r["id_compra_item"]: r for r in ler(f"julgamentos_modelo_f2_modelo_{rod}.csv") if int(r["ordem"]) == o}
    for i in sorted(fichas, key=h)[:15]:
        r = fichas[i]
        linhas.append(["Modelo de aquisição", i, cat_txt(cats_item.get(i)), r["classe"], r["objeto_compra"], r["texto_item"], venc_item.get(i) or "sem resultado", ""])
        ant.append(["Modelo de aquisição", i, r["classe"], f"f2_modelo_{rod}", jul[i]["pertinente"], jul[i]["motivo"]])
fm = ler("fichas_marcas_f2_marca_a1.csv")
jm = {(r["grupo"], r["id_compra"]): r for r in ler("julgamentos_marcas_f2_marca_a1.csv")}
for grupo, n, rot in (("marca", 20, "Marca e tipo de vendedor"), ("empresa publica de TI", 10, "Empresa pública de TI")):
    for r in sorted([x for x in fm if x["grupo"] == grupo], key=lambda x: h(x["id_compra"]))[:n]:
        classe = ("Marcas detectadas e tipo de vendedor (ver colunas F e G)" if grupo == "marca" else "Vencedor classificado como empresa pública de TI")
        linhas.append([rot, r["id_compra"], cat_txt(cats_compra.get(r["id_compra"])), classe, r["objeto"], r["marcas_e_trechos"] or "-", r["vencedores_e_tipos"], ""])
        j = jm[(grupo, r["id_compra"])]
        ant.append([rot, r["id_compra"], classe, "f2_marca_a1", j["pertinente"], j["motivo"]])

wb = Workbook()
ws = wb.active
ws.title = "revisao"
cab = ["grupo", "id", "categoria", "classe atribuída", "texto do objeto", "texto do item (ou marcas e trechos)", "fornecedor", "revisão (certo/errado)"]
ws.append(cab)
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # noqa: E402

for l in linhas:  # caracteres de controle do texto original sao retirados (o Excel nao os aceita)
    ws.append([ILLEGAL_CHARACTERS_RE.sub("", x) if isinstance(x, str) else x for x in l])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1B4A82")
for col, w in zip("ABCDEFGH", (22, 26, 26, 34, 60, 70, 45, 22)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
inst = wb.create_sheet("instrucoes")
for t in ["Revisão humana da classificação da Fase 2 (subamostra dos itens e compras já julgados; semente " + SEM + ").",
          "Para cada linha, escreva 'certo' ou 'errado' na coluna H.",
          "Modelo de aquisição: a classe atribuída está certa quando o texto do item (e, como contexto, o objeto) sustenta a classe e nenhuma classe de precedência mais alta se aplica. "
          "Precedência: 1 Acréscimo a licença existente; 2 Subscrição ou serviço (inclui licenças temporárias, nuvem e certificados digitais); 3 Licenciamento de direitos permanentes de uso; "
          "4 Suporte e manutenção de licença; 5 Desenvolvimento e serviços técnicos (inclui instalação, manutenção e garantia de equipamentos; não inclui treinamento); "
          "6 Equipamentos e infraestrutura física; 7 Não identificado (nenhuma das anteriores, ou texto livre incompatível com o código do catálogo).",
          "Marca: certo quando cada marca detectada é objeto da compra (vendida, licenciada, suportada ou integrada) e o tipo de vendedor está correto; menção incidental (compatibilidade, link, comparação) é errado.",
          "Empresa pública de TI: certo quando o vencedor é empresa estatal cuja atividade é tecnologia da informação.",
          "Detalhes em fase2/METODOLOGIA_fase2.md."]:
    inst.append([t])
inst.column_dimensions["A"].width = 160
for row in inst.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(V + "revisao_humana.xlsx")
with open(V + "revisao_humana_julgamentos_anteriores.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.writer(fo, delimiter=";")
    w.writerow(["grupo", "id", "classe_atribuida", "amostra", "pertinente_julgamento_anterior", "motivo"])
    w.writerows(ant)
print(len(linhas), "linhas na planilha")
