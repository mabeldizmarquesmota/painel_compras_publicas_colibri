"""As 10 maiores compras da classe Capacitacao da aba de IA, para revisao humana, e notas nas fichas 29 e 35 da amostra 2.
Executar a partir de Processado/, depois de p1_06:  python validacao_ia/capacitacao_maiores_valores.py
Entradas: tabelas_painel_v3/ia_base.csv, ia_base_classes.csv, ia_resultados.csv, ia_itens.csv, ia_fornecedores_cnpj.csv;
  scripts_parte0/p0.duckdb (somente leitura: pb e textos dos itens); validacao_ia/revisao_humana_ia.xlsx.
Regras:
  - maiores valores: valor nominal da compra na classe Capacitacao (valor repartido pelos itens), as 10 maiores;
  - planilha: as 10 compras entram como linhas extras, grupo 'Capacitação: maiores valores', com a coluna H vazia; as linhas
    do grupo 'Classes da aba IA' (50 casos ja marcados) nao sao lidas para gravar nem alteradas; antes de gravar, o script
    confere que as 50 marcacoes sao iguais as de validacao_ia/revisao_humana_ia_marcada_v1.xlsx e confere de novo depois;
  - fichas 29 e 35 da amostra 2: nota na coluna 'nota' (o julgamento nao muda).
Saidas: validacao_ia/capacitacao_maiores_valores.csv; validacao_ia/revisao_humana_ia.xlsx (linhas extras);
  validacao_ia/fichas_ia_a2.csv (coluna nota)
"""
import csv

import duckdb
from openpyxl import load_workbook
from openpyxl.styles import Alignment

TP, V = "tabelas_painel_v3/", "validacao_ia/"
ler = lambda n: list(csv.DictReader(open(n, encoding="utf-8"), delimiter=";"))
GRUPO = "Capacitação: maiores valores"
COMP = {"A": "Inteligência artificial (categoria do painel)", "B": "Capacitação em IA (fora das seis categorias)", "C": "Acréscimo de IA a licença (outra categoria)"}
base = {r["id_compra"]: r for r in ler(TP + "ia_base.csv")}
vc = {}
for r in ler(TP + "ia_base_classes.csv"):
    if r["classe"] == "Capacitação":
        x = vc.setdefault(r["id_compra"], [0.0, 0.0])
        x[0] += float(r["valor_nominal"])
        x[1] += float(r["valor_deflacionado"])
top = sorted(vc.items(), key=lambda kv: (-kv[1][0], kv[0]))[:10]
ids = [c for c, _ in top]
nomes = {r["cnpj_fornecedor"]: r["nome"] for r in ler(TP + "ia_fornecedores_cnpj.csv")}
venc = {}
for r in ler(TP + "ia_resultados.csv"):
    if r["id_compra"] in ids:
        d = venc.setdefault(r["id_compra"], {})
        d[r["cnpj_fornecedor"]] = d.get(r["cnpj_fornecedor"], 0) + float(r["valor_nominal"])
itens_ia = {r["id_compra_item"] for r in ler(TP + "ia_itens.csv") if r["id_compra"] in ids and r["casa_termo_ia"] == "sim"}
con = duckdb.connect("scripts_parte0/p0.duckdb", read_only=True)
pb = {r[0]: r[1:] for r in con.sql("SELECT id_compra, ano, orgao, objeto_compra, informacao_complementar FROM pb WHERE id_compra IN (SELECT unnest($l))", params={"l": ids}).fetchall()}
txt = {}
for c, i, t in con.sql("""SELECT id_compra, id_compra_item, concat_ws(' | ', descricao, descricao_detalhada, descricao_resumida) FROM read_parquet('itens_dedup.parquet')
    WHERE id_compra IN (SELECT unnest($l))""", params={"l": ids}).fetchall():
    if i in itens_ia:
        txt.setdefault(c, []).append(t)
con.close()
linhas = []
for k, (c, (vn, vd)) in enumerate(top, 1):
    ano, orgao, obj, info = pb[c]
    v = venc.get(c, {})
    fm = nomes.get(max(v, key=v.get)) if v else "sem resultado com valor"
    linhas.append([k, c, ano, orgao, COMP[base[c]["componente"]], round(vn, 2), round(vd, 2), fm, (obj or "").strip(), (info or "").strip()[:500], " || ".join(txt.get(c, []))[:1500]])
with open(V + "capacitacao_maiores_valores.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.writer(fo, delimiter=";")
    w.writerow(["posicao", "id_compra", "ano", "orgao", "componente", "valor_nominal", "valor_deflacionado", "fornecedor_principal", "objeto", "informacao_complementar", "itens_com_termo_de_ia"])
    w.writerows(linhas)


def marcas(arq):
    ws = load_workbook(arq)["revisao"]
    return [(r[1], r[7]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "Classes da aba IA"]


ref = marcas(V + "revisao_humana_ia_marcada_v1.xlsx")
assert len(ref) == 50 and marcas(V + "revisao_humana_ia.xlsx") == ref, "as 50 marcações mudaram antes da gravação"
wb = load_workbook(V + "revisao_humana_ia.xlsx")
ws = wb["revisao"]
# a lista revisada fica congelada: se alguma linha extra ja tem marcacao, a planilha nao e regravada
if any(ws.cell(i, 1).value == GRUPO and ws.cell(i, 8).value not in (None, "") for i in range(2, ws.max_row + 1)):
    raise SystemExit("linhas extras já revisadas: a planilha não é regravada (lista em validacao_ia/capacitacao_maiores_valores_revisadas.csv)")
# remove so as linhas extras de uma execucao anterior (de baixo para cima), nunca as 50 marcadas
for i in range(ws.max_row, 1, -1):
    if ws.cell(i, 1).value == GRUPO:
        ws.delete_rows(i)
for l in linhas:
    ws.append([GRUPO, l[1], l[4], "Capacitação", (l[8] + (" | " + l[9] if l[9] else ""))[:1500], l[10] or "(compra classificada pelo objeto)", l[7], None])
for row in ws.iter_rows(min_row=ws.max_row - len(linhas) + 1):
    for cel in row:
        cel.alignment = Alignment(wrap_text=True, vertical="top")
wi = wb["instrucoes"]
NOTA_INSTR = ("Linhas extras do grupo 'Capacitação: maiores valores' (as 10 compras de maior valor na classe Capacitação): julgue com o mesmo critério da classe Capacitação; "
              "elas não entram no cálculo da concordância das 50 linhas já revisadas.")
if not any(r[0] == NOTA_INSTR for r in wi.iter_rows(values_only=True)):
    wi.append([NOTA_INSTR])
    wi.cell(wi.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
wb.save(V + "revisao_humana_ia.xlsx")
assert marcas(V + "revisao_humana_ia.xlsx") == ref, "as 50 marcações mudaram"
# nota nas fichas 29 e 35 da amostra 2 (classe Acrescimo): o julgamento nao muda
NOTAS = {"29": "Com a repartição do valor pelos itens (regra do gráfico 9), a compra passou a ser contada em Solução de IA (exclusiva ou em pacote), porque os itens de Copilot valem menos que o restante; o julgamento não muda.",
         "35": "Com a repartição do valor pelos itens (regra do gráfico 9), a compra passou a ser contada em Solução de IA (exclusiva ou em pacote), porque os itens de Copilot valem menos que o restante; o julgamento não muda."}
fa2 = ler(V + "fichas_ia_a2.csv")
cab = list(fa2[0].keys()) + ([] if "nota" in fa2[0] else ["nota"])
for r in fa2:
    r["nota"] = NOTAS.get(r["ficha"], r.get("nota", ""))
    assert r["ficha"] not in NOTAS or base[r["id_compra"]]["classe"] == "Solução de IA (exclusiva ou em pacote)"
with open(V + "fichas_ia_a2.csv", "w", encoding="utf-8", newline="") as fo:
    w = csv.DictWriter(fo, fieldnames=cab, delimiter=";")
    w.writeheader()
    w.writerows(fa2)
for l in linhas:
    print(l[0], l[1], l[2], l[3][:45], "|", round(l[5] / 1e6, 2), "mi |", l[7][:45], "|", l[8][:90])
print("50 marcações iguais à cópia marcada_v1: sim")
